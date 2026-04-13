"""
Excel/CSV importer for marketing client lists.
"""
import csv
import io
import re
from typing import Any

from openpyxl import load_workbook

# Column aliases — maps raw header names to internal field names
_COLUMN_ALIASES: dict[str, str] = {
    "name": "name",
    "nama": "name",
    "nama perusahaan": "name",
    "company": "name",
    "company name": "name",
    "organization": "name",
    "organisation": "name",
    "perusahaan": "name",
    "client": "name",
    "client name": "name",
    "lembaga": "name",
    "instansi": "name",
    # Province
    "province": "province",
    "provinsi": "province",
    # Website
    "website": "website",
    "web": "website",
    "url": "website",
    # Email
    "email": "email",
    "e-mail": "email",
    "email address": "email",
    "alamat email": "email",
    "email kampus": "email",
    "email kantor": "email",
    "official email": "email",
    "mail": "email",
    # Phone
    "phone": "phone",
    "telephone": "phone",
    "telp": "phone",
    "no telepon": "phone",
    # Contact Person
    "contact person": "contact_person",
    "pic": "contact_person",
    "cp": "contact_person",
}


def _canonicalize_header(value: str) -> str:
    normalized = re.sub(r"[_\-.]+", " ", str(value or "").strip().lower())
    return re.sub(r"\s+", " ", normalized).strip()


def _looks_like_email(value: Any) -> bool:
    if value is None:
        return False
    text = str(value).strip()
    return bool(re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", text))


def _infer_field_from_header(header: str) -> str | None:
    key = _canonicalize_header(header)
    if not key:
        return None

    if key in _COLUMN_ALIASES:
        return _COLUMN_ALIASES[key]

    if "email" in key or key == "mail":
        return "email"
    if "phone" in key or "telepon" in key or "telp" in key:
        return "phone"
    if "province" in key or "provinsi" in key:
        return "province"
    if "website" in key or key in {"web", "url"}:
        return "website"
    if "contact person" in key or ("contact" in key and "name" in key):
        return "contact_person"
    if key in {"nama", "name"} or "company" in key or "organization" in key or "organisation" in key:
        return "name"

    return None


def _normalize_headers(raw_headers: list[str]) -> tuple[dict[int, str], bool]:
    """Map column indices to internal field names using aliases and heuristics."""
    mapping: dict[int, str] = {}
    found_name = False
    found_email = False

    for i, header in enumerate(raw_headers):
        field = _infer_field_from_header(header)
        if field and field not in mapping.values():
            mapping[i] = field
            if field == "name":
                found_name = True
            if field == "email":
                found_email = True

    if not found_name and raw_headers:
        mapping[0] = "name"

    return mapping, found_email


def _infer_mapping_from_first_row(first_row: list[Any]) -> dict[int, str]:
    """Fallback for files where the first row is already data, not a header."""
    mapping: dict[int, str] = {}
    email_idx: int | None = None

    for i, value in enumerate(first_row):
        if _looks_like_email(value):
            email_idx = i
            mapping[i] = "email"
            break

    if email_idx is None:
        return mapping

    for i, value in enumerate(first_row):
        if i == email_idx:
            continue
        if value is not None and str(value).strip():
            mapping[i] = "name"
            break

    return mapping


def _rows_from_worksheet(ws, include_first_row_as_data: bool = False, first_row: tuple[Any, ...] | None = None):
    if include_first_row_as_data and first_row is not None:
        yield first_row
    for row in ws.iter_rows(min_row=2, values_only=True):
        yield row


def _parse_tabular_rows(rows: list[list[Any]], raw_headers: list[str], header_is_data: bool = False) -> list[dict[str, Any]]:
    rows_data: list[dict[str, Any]] = []

    if header_is_data:
        col_map = _infer_mapping_from_first_row(rows[0]) if rows else {}
        effective_headers = [f"Column {i + 1}" for i in range(len(raw_headers))]
        iterable_rows = rows
    else:
        col_map, _ = _normalize_headers(raw_headers)
        effective_headers = raw_headers
        iterable_rows = rows

    for row in iterable_rows:
        row_dict: dict[str, Any] = {}
        extra: dict[str, Any] = {}
        for i, val in enumerate(row):
            if i in col_map:
                row_dict[col_map[i]] = str(val).strip() if val is not None else ""
            elif val is not None and str(val).strip():
                header_name = effective_headers[i] if i < len(effective_headers) else f"Column {i + 1}"
                extra[str(header_name)] = str(val).strip()
        if extra:
            row_dict["_extra"] = extra
        rows_data.append(row_dict)

    return rows_data


def parse_excel_file(file_path: str) -> tuple[list[dict[str, Any]], list[str]]:
    """Parse an Excel or CSV file and return (rows, detected_columns)."""
    ext = file_path.lower().split(".")[-1]

    if ext in ("xlsx", "xls"):
        wb = load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            wb.close()
            return [], []

        raw_headers = [str(h).strip() if h else "" for h in header_row]
        _, found_email = _normalize_headers(raw_headers)
        header_is_data = not found_email and any(_looks_like_email(value) for value in header_row)
        rows = [list(row) for row in _rows_from_worksheet(ws, include_first_row_as_data=header_is_data, first_row=header_row)]
        wb.close()
        return _parse_tabular_rows(rows, raw_headers, header_is_data=header_is_data), raw_headers

    if ext == "csv":
        with open(file_path, newline="", encoding="utf-8") as f:
            reader = csv.reader(f)
            all_rows = list(reader)
        if not all_rows:
            return [], []

        raw_headers = [str(h).strip() for h in all_rows[0]]
        _, found_email = _normalize_headers(raw_headers)
        header_is_data = not found_email and any(_looks_like_email(value) for value in all_rows[0])
        data_rows = all_rows if header_is_data else all_rows[1:]
        return _parse_tabular_rows(data_rows, raw_headers, header_is_data=header_is_data), raw_headers

    return [], []


def parse_excel_bytes(content: bytes, filename: str) -> tuple[list[dict[str, Any]], list[str]]:
    """Parse Excel/CSV content from bytes. Returns (rows, detected_columns)."""
    ext = filename.lower().split(".")[-1]

    if ext in ("xlsx", "xls"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            wb.close()
            return [], []

        raw_headers = [str(h).strip() if h else "" for h in header_row]
        _, found_email = _normalize_headers(raw_headers)
        header_is_data = not found_email and any(_looks_like_email(value) for value in header_row)
        rows = [list(row) for row in _rows_from_worksheet(ws, include_first_row_as_data=header_is_data, first_row=header_row)]
        wb.close()
        return _parse_tabular_rows(rows, raw_headers, header_is_data=header_is_data), raw_headers

    if ext == "csv":
        text = content.decode("utf-8")
        reader = csv.reader(io.StringIO(text))
        all_rows = list(reader)
        if not all_rows:
            return [], []

        raw_headers = [str(h).strip() for h in all_rows[0]]
        _, found_email = _normalize_headers(raw_headers)
        header_is_data = not found_email and any(_looks_like_email(value) for value in all_rows[0])
        data_rows = all_rows if header_is_data else all_rows[1:]
        return _parse_tabular_rows(data_rows, raw_headers, header_is_data=header_is_data), raw_headers

    return [], []
