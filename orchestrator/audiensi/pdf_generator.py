"""
PDF/Word document generator for audiensi surat undangan.

Uses python-docx to fill in a Word template with {{placeholders}},
saves the filled document. PDF conversion can be done externally
(e.g. LibreOffice headless) if needed.
"""

import os
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any

from orchestrator.config import log, cfg

WIB = timezone(timedelta(hours=7))

# Default template location
DEFAULT_TEMPLATE_DIR = Path("data/templates")
DEFAULT_TEMPLATE_NAME = "surat_undangan.docx"

# Output directory for generated documents
OUTPUT_DIR = Path("data/audiensi_docs")

# Available placeholders with descriptions
PLACEHOLDERS = {
    "nama_universitas": "Nama universitas (e.g. 'Universitas Prisma')",
    "nama_rektor": "Nama rektor atau 'Rektor' jika belum diketahui",
    "tanggal_surat": "Tanggal surat (e.g. '20 Februari 2026')",
    "provinsi": "Provinsi universitas",
    "nomor_surat": "Nomor surat otomatis (e.g. '001/AAI/AUD/II/2026')",
}

# Indonesian month names
_BULAN = [
    "", "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember",
]


def _format_tanggal(dt: datetime) -> str:
    """Format datetime as Indonesian date string."""
    return f"{dt.day} {_BULAN[dt.month]} {dt.year}"


def _generate_nomor_surat(audiensi_id: int) -> str:
    """Generate surat number based on audiensi ID and current date."""
    now = datetime.now(WIB)
    # Roman numeral month
    roman = ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]
    month_roman = roman[now.month]
    return f"{audiensi_id:03d}/AAI/AUD/{month_roman}/{now.year}"


def get_template_path() -> Path:
    """Return the path to the active Word template."""
    custom = cfg.AUDIENSI_PDF_TEMPLATE_PATH if hasattr(cfg, 'AUDIENSI_PDF_TEMPLATE_PATH') else ""
    if custom and Path(custom).exists():
        return Path(custom)
    return DEFAULT_TEMPLATE_DIR / DEFAULT_TEMPLATE_NAME


def template_exists() -> bool:
    """Check if a Word template is available."""
    return get_template_path().exists()


async def generate_audiensi_document(
    audiensi_id: int,
    university_name: str,
    rector_name: str | None = None,
    province: str | None = None,
) -> str | None:
    """
    Generate a filled Word document from the template.

    1. Load Word template from data/templates/surat_undangan.docx
    2. Replace all {{placeholders}} with actual values
    3. Save filled .docx to data/audiensi_docs/{audiensi_id}_surat.docx
    4. Return file path or None on error

    For now, returns .docx path. PDF conversion can be added later
    via LibreOffice headless or docx2pdf.
    """
    template_path = get_template_path()
    if not template_path.exists():
        log.warning("Audiensi template not found at %s", template_path)
        return None

    try:
        from docx import Document
    except ImportError:
        log.error("python-docx not installed. Run: pip install python-docx")
        return None

    try:
        doc = Document(str(template_path))

        # Prepare replacement values
        now = datetime.now(WIB)
        replacements = {
            "{{nama_universitas}}": university_name or "",
            "{{nama_rektor}}": rector_name or "Rektor",
            "{{tanggal_surat}}": _format_tanggal(now),
            "{{provinsi}}": province or "",
            "{{nomor_surat}}": _generate_nomor_surat(audiensi_id),
        }

        # Replace placeholders in all paragraphs
        for paragraph in doc.paragraphs:
            _replace_in_paragraph(paragraph, replacements)

        # Replace placeholders in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        _replace_in_paragraph(paragraph, replacements)

        # Replace in headers/footers
        for section in doc.sections:
            for header_footer in [section.header, section.footer]:
                if header_footer:
                    for paragraph in header_footer.paragraphs:
                        _replace_in_paragraph(paragraph, replacements)

        # Ensure output directory exists
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

        output_path = OUTPUT_DIR / f"{audiensi_id}_surat.docx"
        doc.save(str(output_path))

        log.info("Generated audiensi document: %s", output_path)
        return str(output_path)

    except Exception as e:
        log.error("Failed to generate audiensi document for ID %d: %s", audiensi_id, e)
        return None


def _replace_in_paragraph(paragraph, replacements: dict[str, str]) -> None:
    """
    Replace {{placeholders}} in a paragraph while preserving formatting.

    python-docx splits text across multiple runs, so we need to handle
    the case where a placeholder spans multiple runs.
    """
    # First try simple run-by-run replacement
    for run in paragraph.runs:
        for placeholder, value in replacements.items():
            if placeholder in run.text:
                run.text = run.text.replace(placeholder, value)

    # Handle placeholders split across runs
    full_text = paragraph.text
    for placeholder, value in replacements.items():
        if placeholder in full_text:
            # Rebuild the paragraph text preserving the first run's format
            _replace_across_runs(paragraph, placeholder, value)


def _replace_across_runs(paragraph, placeholder: str, value: str) -> None:
    """Replace a placeholder that may span multiple runs."""
    runs = paragraph.runs
    if not runs:
        return

    # Build a map of character positions to runs
    combined = ""
    for run in runs:
        combined += run.text

    if placeholder not in combined:
        return

    # Find where the placeholder starts and ends
    start_idx = combined.index(placeholder)
    end_idx = start_idx + len(placeholder)

    # Find which runs contain the placeholder
    char_count = 0
    start_run = end_run = 0
    start_offset = end_offset = 0

    for i, run in enumerate(runs):
        run_start = char_count
        run_end = char_count + len(run.text)

        if run_start <= start_idx < run_end:
            start_run = i
            start_offset = start_idx - run_start
        if run_start < end_idx <= run_end:
            end_run = i
            end_offset = end_idx - run_start
            break

        char_count = run_end

    # Replace text in affected runs
    if start_run == end_run:
        # Placeholder is within a single run
        run = runs[start_run]
        run.text = run.text[:start_offset] + value + run.text[end_offset:]
    else:
        # Placeholder spans multiple runs
        runs[start_run].text = runs[start_run].text[:start_offset] + value
        for i in range(start_run + 1, end_run):
            runs[i].text = ""
        runs[end_run].text = runs[end_run].text[end_offset:]


async def save_uploaded_template(file_content: bytes, filename: str) -> str:
    """Save an uploaded template file."""
    DEFAULT_TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    dest = DEFAULT_TEMPLATE_DIR / DEFAULT_TEMPLATE_NAME
    dest.write_bytes(file_content)
    log.info("Saved audiensi template: %s (%d bytes)", dest, len(file_content))
    return str(dest)


def get_available_placeholders() -> list[dict[str, str]]:
    """Return list of available placeholders with descriptions."""
    return [{"key": f"{{{{{k}}}}}", "description": v} for k, v in PLACEHOLDERS.items()]
